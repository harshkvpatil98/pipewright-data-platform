"""Chart shaping, the exported file, and finding things in the catalog."""

from __future__ import annotations

import io
import zipfile

import pandas as pd
import pytest

from service_reporting.aggregation import Measure, Query, run_query
from service_reporting.catalog import SearchableDataset, search, slugify, tokenize
from service_reporting.charts import CHART_TYPES_BY_NAME, to_chart_data, validate_chart
from service_reporting.exports import export, safe_filename
from shared_python.errors import BadRequestError

FRAME = pd.DataFrame(
    {
        "region": ["north", "south", "east", "west"],
        "month": ["jan", "jan", "feb", "feb"],
        "amount": [100, 250, 120, 80],
    }
)


def _query(**overrides) -> Query:
    base = {"dimensions": ["region"], "measures": [Measure("amount", "sum", "total")]}
    base.update(overrides)
    return Query(**base)


# ---- chart validation ----


def test_a_pie_chart_refuses_two_grouping_columns():
    with pytest.raises(BadRequestError) as caught:
        validate_chart("pie", _query(dimensions=["region", "month"]))
    assert "at most 1 grouping" in str(caught.value.detail)


def test_a_kpi_refuses_a_grouping_column_because_it_is_one_number():
    with pytest.raises(BadRequestError):
        validate_chart("kpi", _query(dimensions=["region"]))


def test_a_scatter_needs_two_measures():
    with pytest.raises(BadRequestError) as caught:
        validate_chart("scatter", _query(dimensions=[]))
    assert "at least 2 measure" in str(caught.value.detail)


def test_an_unknown_chart_type_lists_the_ones_that_exist():
    with pytest.raises(BadRequestError) as caught:
        validate_chart("sankey", _query())
    assert "bar" in str(caught.value.detail)


def test_a_limit_beyond_what_a_pie_can_show_is_warned_about():
    warnings = validate_chart("pie", _query(limit=100))
    assert any("stops being readable" in warning for warning in warnings)


# ---- chart shaping ----


def test_a_bar_chart_becomes_labels_and_one_series():
    query = _query()
    data = to_chart_data("bar", query, run_query(FRAME, query))
    assert set(data.labels) == {"north", "south", "east", "west"}
    assert [series["name"] for series in data.series] == ["total"]


def test_a_kpi_is_a_single_value():
    query = _query(dimensions=[])
    data = to_chart_data("kpi", query, run_query(FRAME, query))
    assert data.series[0]["values"] == [550]
    assert data.labels == []


def test_a_second_dimension_becomes_one_series_per_value():
    """'Revenue by month, split by region' is two dimensions, not two charts."""
    query = Query(
        dimensions=["month", "region"], measures=[Measure("amount", "sum", "total")]
    )
    data = to_chart_data("line", query, run_query(FRAME, query))
    assert len(data.series) == 4  # one per region


def test_a_pie_chart_is_capped_at_what_it_can_show():
    wide = pd.DataFrame({"k": [f"k{i}" for i in range(40)], "v": range(40)})
    query = Query(dimensions=["k"], measures=[Measure("v", "sum", "total")])
    data = to_chart_data("pie", query, run_query(wide, query))
    assert len(data.labels) == CHART_TYPES_BY_NAME["pie"].max_categories
    assert any("cannot show more legibly" in warning for warning in data.warnings)


def test_an_empty_result_produces_an_empty_chart_not_an_error():
    empty = FRAME.head(0)
    query = _query()
    data = to_chart_data("bar", query, run_query(empty, query))
    assert data.labels == []
    assert data.row_count == 0


def test_nulls_and_infinities_never_reach_the_json():
    frame = pd.DataFrame({"k": ["a", "b"], "v": [1.0, float("inf")]})
    query = Query(dimensions=["k"], measures=[Measure("v", "max", "peak")])
    data = to_chart_data("bar", query, run_query(frame, query))
    values = data.series[0]["values"]
    assert all(value is None or isinstance(value, (int, float)) for value in values)
    assert not any(isinstance(value, float) and value == float("inf") for value in values)


# ---- exports ----


def test_an_excel_export_is_a_real_workbook():
    result = export(FRAME, title="Revenue", file_format="excel")
    # A .xlsx is a zip; a CSV with the wrong extension is not.
    with zipfile.ZipFile(io.BytesIO(result.content)) as archive:
        assert any(name.endswith("workbook.xml") for name in archive.namelist())


def test_the_excel_header_row_is_frozen_so_a_long_report_stays_readable():
    from openpyxl import load_workbook

    result = export(FRAME, title="Revenue", file_format="excel", subtitle="Last month")
    sheet = load_workbook(io.BytesIO(result.content)).active
    assert sheet.freeze_panes is not None
    assert sheet.auto_filter.ref is not None


def test_numeric_columns_get_a_number_format_not_text():
    from openpyxl import load_workbook

    result = export(FRAME, title="Revenue", file_format="excel")
    sheet = load_workbook(io.BytesIO(result.content)).active
    amount_column = [cell for cell in sheet[sheet.max_row] if cell.number_format != "General"]
    assert amount_column


def test_an_html_export_is_self_contained():
    """No stylesheet to lose, and it prints to PDF correctly from a browser."""
    result = export(FRAME, title="Revenue", file_format="html")
    text = result.content.decode()
    assert "<style>" in text
    assert "http://" not in text and "https://" not in text


def test_pdf_is_a_real_file_and_an_unknown_format_is_still_refused():
    # This test used to pin the honest refusal "no layout engine to render one
    # properly". P8 added reportlab, a maintained pure-Python layout engine, so
    # the refusal is gone and the format is real; the guard now protects the
    # format list itself.
    pdf = export(FRAME, title="Revenue", file_format="pdf")
    assert pdf.content[:5] == b"%PDF-" and pdf.media_type == "application/pdf"
    with pytest.raises(BadRequestError) as caught:
        export(FRAME, title="Revenue", file_format="docx")
    assert "pdf" in str(caught.value.detail)


def test_a_filename_survives_every_filesystem():
    name = safe_filename("Q3 revenue / EMEA: draft", "excel")
    assert "/" not in name and ":" not in name
    assert name.endswith(".xlsx")


def test_an_enormous_export_is_refused_with_advice():
    huge = pd.DataFrame({"id": range(200_001)})
    with pytest.raises(BadRequestError) as caught:
        export(huge, title="Everything", file_format="csv")
    assert "Filter it down" in str(caught.value.detail)


# ---- catalog ----


def _datasets() -> list[SearchableDataset]:
    return [
        SearchableDataset(
            id="1", name="orders", columns=["id", "customer_id", "amount"],
            certified=True, tags=["finance"],
        ),
        SearchableDataset(
            id="2", name="order_lines", columns=["order_id", "sku"],
            description="Lines within each order",
        ),
        SearchableDataset(id="3", name="customers", columns=["id", "email", "order_count"]),
        SearchableDataset(id="4", name="order_line_adjustments", columns=["order_id"]),
    ]


def test_the_closest_name_match_wins_over_incidental_column_hits():
    """Searching 'order' must find 'orders', not 'order_lines'."""
    hits = search(_datasets(), "order")
    assert [hit.name for hit in hits][:2] == ["orders", "order_lines"]


def test_a_shorter_name_beats_a_longer_one_for_the_same_prefix():
    hits = search(_datasets(), "order")
    names = [hit.name for hit in hits]
    assert names.index("order_lines") < names.index("order_line_adjustments")


def test_a_column_match_still_finds_a_dataset_nobody_named_well():
    hits = search(_datasets(), "email")
    assert [hit.name for hit in hits] == ["customers"]


def test_search_says_why_it_matched():
    hit = next(hit for hit in search(_datasets(), "finance"))
    assert any("tagged" in reason for reason in hit.reasons)


def test_certified_datasets_win_ties():
    hits = search(_datasets(), "orders")
    assert hits[0].certified is True


def test_an_empty_query_lists_everything_rather_than_nothing():
    """A catalog you can only use by knowing what to type is not a catalog."""
    hits = search(_datasets(), "")
    assert len(hits) == 4
    assert hits[0].name == "orders"  # certified first


def test_filters_narrow_the_catalog():
    assert len(search(_datasets(), "", certified_only=True)) == 1
    assert len(search(_datasets(), "", tag="finance")) == 1
    assert len(search(_datasets(), "", tag="nonexistent")) == 0


def test_single_characters_are_not_worth_searching_for():
    assert tokenize("a b orders") == ["orders"]


def test_a_search_matching_nothing_returns_nothing():
    assert search(_datasets(), "quantum") == []


def test_glossary_slugs_are_stable_and_url_safe():
    assert slugify("Monthly Recurring Revenue") == "monthly-recurring-revenue"
    assert slugify("  Gross Margin %  ") == "gross-margin"
    assert slugify("!!!") == "term"
